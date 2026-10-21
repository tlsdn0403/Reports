from pathlib import Path
from docx import Document
from docx.shared import Pt,Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from zipfile import ZipFile, ZIP_DEFLATED
from io import BytesIO
src=Path(r'C:/Users/tlsdn/OneDrive/문서/카카오톡 받은 파일/이력서 자소서 양식_260908.docx')
d=Document(src)
from copy import deepcopy
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
# Reuse the template's populated-row design for a third record in spacer slots.
for target,source in [(9,8),(24,23)]:
 table=d.tables[1]; old=table.rows[target]._tr; old.getparent().replace(old,deepcopy(table.rows[source]._tr))
def fill(t,r,c,text,size=9):
 cell=d.tables[t].rows[r].cells[c]
 p=cell.paragraphs[0]
 for other in list(cell.paragraphs)[1:]: other._element.getparent().remove(other._element)
 p.clear(); run=p.add_run(text); run.font.size=Pt(size)
 return p
for r,c,s in [(0,3,'박신우'),(0,5,'Park Shin Woo'),(0,7,'게임 클라이언트'),(1,3,'2001.04.03'),(1,7,'만 25세'),(3,3,'tlsdn0403@gmail.com'),(4,3,'경기도 용인시 수지구 정평로 73')]: fill(0,r,c,s)
p=fill(0,0,0,''); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p.add_run().add_picture('제출용 자료/증명사진.jpg',width=Cm(2.6))
for c,s in [(0,'2021.03~2027.02'),(3,'한국공학대학교 게임공학과'),(4,'3.5 / 4.5'),(6,'졸업 예정')]: fill(1,2,c,s)
for r,values in [(7,['2025.12~2026.07','감염: 죽음의 도시 · 절단 및 트럭 탑승','졸업작품','한국공학대학교']),(8,['2025.11~2025.12','NGP · 캐릭터 송수신 및 충돌 처리','전공 프로젝트','한국공학대학교']),(9,['2026.07~2026.08','FlickDom · 보드 판정 및 턴 동기화','팀 프로젝트','3인 팀'])]:
 for c,s in zip([0,3,4,6],values): fill(1,r,c,s)
for c,s in [(0,'OPIc'),(3,'IM2'),(4,'2023.10')]: fill(1,12,c,s)
for r,values in [(22,['Unreal Engine 5','C++ 게임플레이 구현','좀비 절단·트럭 탑승 구현']),(23,['Unity 6','C# 보드게임 개발','보드 판정·턴 상태 동기화']),(24,['Visual Studio / Git','C++ 개발·버전 관리','디버깅·팀 코드 통합'])]:
 for c,s in zip([0,3,5],values): fill(1,r,c,s)
fill(1,26,4,'계급:'); fill(1,26,5,'전역(해제)사유: 군필')
answers=[
'졸업작품 「감염: 죽음의 도시」에서 UE5와 C++로 좀비 신체 절단과 트럭 탑승 기능을 맡았습니다. 트럭 위의 캐릭터가 튀어 오르거나 뒤로 밀리는 문제가 있었습니다. 먼저 차량과 캐릭터의 충돌을 분리하고, 트럭 바닥을 이동 기준으로 지정하는 SetBase를 적용해 차량을 따라가도록 했습니다. 이후 하차 시 위치가 순간적으로 보정되는 현상이 남아, SetBase(nullptr)로 이동 기준을 해제하고 관련 상태도 초기화했습니다. 이를 통해 충돌 처리와 함께 탑승·하차의 상태 전환도 확인해야 한다는 점을 배웠습니다.',
'필요한 기능을 정한 뒤 관련 코드와 엔진 문서를 찾아보고, 프로젝트에 적용하면서 학습합니다. 트럭 탑승 기능에서는 이동 기준과 좌표 변환을 공부해, 적재함 내 위치를 차량 기준 좌표로 전송하고 수신 측에서 월드 좌표로 복원했습니다. 최근에는 Codex와 Unity MCP를 활용해 Unity 프로젝트 FlickDom을 개발했습니다. 보드 점유와 카드 패턴 판정, 턴 상태 동기화를 맡았으며, 이전 턴의 데이터가 뒤늦게 적용되지 않도록 라운드와 턴 정보를 비교하는 처리를 추가했습니다.',
'팀 프로젝트에서 선행 리팩터링이 늦어지면서 후속 작업을 맡은 팀원과 의견이 엇갈렸습니다. 소통까지 줄어들자 저는 두 팀원이 함께 이야기할 자리를 마련하고, 각자 어려운 점과 상대에게 바라는 점을 들었습니다. 누가 잘못했는지 판단하기보다 프로젝트를 함께 완성해야 한다는 목표를 중심으로 대화를 이어갔고, 이후 소통이 회복되어 작업을 계속할 수 있었습니다. 이 경험을 통해 의견이 다를 때 피드백을 미루기보다 상대의 상황을 먼저 듣고, 작업에 미치는 영향을 구체적으로 이야기해야 한다는 점을 배웠습니다.',
'C++로 캐릭터 상태를 처리하고 화면의 움직임으로 연결할 수 있습니다. NGP 프로젝트에서는 동기화에 필요한 정보를 먼저 정했습니다. 캐릭터 ID, 위치, 이동 방향, 회전각, 팔다리 움직임 각도를 구조체에 담아 서버로 보내고, 수신한 정보를 다른 플레이어의 캐릭터 객체에 반영했습니다. 저는 캐릭터 정보 송수신과 적용, 충돌 처리를 담당했으며 서버 연결과 스레드 구성은 팀원이 맡았습니다. UE5에서는 부위별 내구도가 소진되면 해당 본을 숨기고 물리 효과를 적용한 절단 조각을 생성하는 기능도 구현했습니다.'
]
for r,s in zip([1,3,5,7],answers):
 p=fill(2,r,0,s,10); p.paragraph_format.line_spacing=Pt(14); p.paragraph_format.space_after=Pt(0)
 p.alignment=WD_ALIGN_PARAGRAPH.LEFT
 pr=p._p.get_or_add_pPr()
 for old in list(pr.findall(qn('w:snapToGrid'))):pr.remove(old)
 grid=OxmlElement('w:snapToGrid');grid.set(qn('w:val'),'0');pr.append(grid)
for p in d.paragraphs:
 if '성  명' in p.text or '성  명' in p.text.replace('\t',''):
  p.text='\t\t\t\t\t\t\t\t성  명 : 박신우      (인)'
# Recover vertical room from unused blank education/language rows.
for ri in [3,13]:
 h=d.tables[1].rows[ri]._tr.find(qn('w:trPr')).find(qn('w:trHeight'));h.set(qn('w:val'),'180')
buf=BytesIO(); d.save(buf); buf.seek(0)
out=Path('output/박신우_이력서_자기소개서_기본작성본.docx');out.parent.mkdir(exist_ok=True)
with ZipFile(src) as original,ZipFile(buf) as edited,ZipFile(out,'w',ZIP_DEFLATED) as final:
 for n in original.namelist():
  final.writestr(n, edited.read(n) if n in ['word/document.xml','word/_rels/document.xml.rels','[Content_Types].xml'] else original.read(n))
 for n in edited.namelist():
  if n not in original.namelist():final.writestr(n,edited.read(n))
print(out.resolve())
