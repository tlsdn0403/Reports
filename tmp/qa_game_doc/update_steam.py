from pathlib import Path
from PIL import Image
from docx import Document
from docx.shared import Inches,Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import numpy as np
base=Path('회사지원/펄어비스');src=base/'steamcommunity.com-profiles-76561198149522185-games.png'
im=Image.open(src);w,h=im.size
# Page crops preserve the original embedded image and its aspect ratio.
a=np.asarray(im.convert('RGB'))
maxh=int(w*8.55/6.5);bounds=[0]
while h-bounds[-1]>maxh:
 target=bounds[-1]+maxh
 lo=target-350;hi=target
 # Broad, flat rows in the content area are gaps between game entries.
 v=a[lo:hi,int(w*.11):int(w*.90),:].astype(float)
 scores=np.abs(v-a[1000,20].astype(float)).mean(axis=(1,2))
 valid=np.where(scores<1)[0]; cut=lo+int(valid[-1] if len(valid) else scores.argmin())
 bounds.append(cut)
bounds.append(h)
d=Document(base/'박신우_게임플레이_경험_사진추가용.docx')
pars=d.paragraphs
start=next(p for p in pars if p.text=='Steam 라이브러리 기록')
end=next(p for p in pars if p.text=='경쟁 게임 경험')
# Remove the old screenshot page body but retain the next page break.
node=start._p.getnext()
while node is not end._p:
 nxt=node.getnext()
 if node.xpath('.//w:br[@w:type="page"]'):break
 node.getparent().remove(node);node=nxt
anchor=start._p

def crop_pic(p,y0,y1):
 pic=p.add_run().add_picture(str(src),width=Inches(6.5),height=Inches(6.5*(y1-y0)/w))
 fill=pic._inline.graphic.graphicData.pic.blipFill
 rect=OxmlElement('a:srcRect');rect.set('t',str(round(y0/h*100000)));rect.set('b',str(round((h-y1)/h*100000)));fill.insert(1,rect)
 pic._inline.docPr.set('descr',f'Steam 게임 목록 캡처 구간 {y0}부터 {y1}')
 p.paragraph_format.space_after=Pt(2)
 return p
p=d.add_paragraph('계정 tlsdn0403  |  플레이 시간순 목록  |  이어지는 기록은 문서 뒤쪽에 수록')
p.paragraph_format.space_after=Pt(6);anchor.addnext(p._p);anchor=p._p
p=crop_pic(d.add_paragraph(),bounds[0],bounds[1]);anchor.addnext(p._p)
for i in range(1,len(bounds)-1):
 d.add_page_break();d.add_heading(f'Steam 라이브러리 기록 {i+1}',1)
 p=d.add_paragraph('계정 tlsdn0403  |  플레이 시간순 목록 이어보기');p.paragraph_format.space_after=Pt(6)
 crop_pic(d.add_paragraph(),bounds[i],bounds[i+1])
for part in [d.styles.element,d.element]:
 for el in list(part.iter()):
  if el.tag.endswith('}pBdr'):el.getparent().remove(el)
out=base/'박신우_게임플레이_경험_스팀증빙반영.docx';d.save(out)
print('size',w,h,'bounds',bounds,'screenshot pages',len(bounds)-1)
