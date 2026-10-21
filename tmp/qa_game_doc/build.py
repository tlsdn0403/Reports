from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from PIL import Image, ImageDraw, ImageFont
root=Path.cwd(); work=root/'tmp/qa_game_doc'; out=root/'회사지원/펄어비스'; out.mkdir(parents=True,exist_ok=True)
doc=Document(); sec=doc.sections[0]; sec.page_height=Inches(11.69); sec.page_width=Inches(8.27)
sec.top_margin=sec.bottom_margin=Inches(.65); sec.left_margin=sec.right_margin=Inches(.75)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
 s=doc.styles[name]; s.font.name='맑은 고딕'; s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'맑은 고딕'); s.font.color.rgb=RGBColor(0,0,0)
 s.paragraph_format.space_after=Pt(7)
doc.styles['Normal'].font.size=Pt(10)
doc.styles['Normal'].paragraph_format.line_spacing=1.15
doc.styles['Title'].font.size=Pt(25)
doc.styles['Heading 1'].font.size=Pt(18)
doc.styles['Heading 2'].font.size=Pt(12)
doc.core_properties.author='박신우';doc.core_properties.title='박신우 게임 플레이 경험'
def p(t='',size=None,color=None):
 x=doc.add_paragraph(t)
 if size:
  for r in x.runs:r.font.size=Pt(size)
 if color:
  for r in x.runs:r.font.color.rgb=RGBColor.from_string(color)
 return x
def heading(t):doc.add_heading(t,level=1)
def new(t):doc.add_page_break();heading(t)
def slot(label,height=1.85):
 idx=len(list(work.glob('slot*.png')))+1
 im=Image.new('RGB',(1600,int(1600*height/6.5)),(247,249,251));dr=ImageDraw.Draw(im)
 dr.rectangle((2,2,im.width-3,im.height-3),outline=(185,195,205),width=3)
 f=ImageFont.truetype('C:/Windows/Fonts/malgun.ttf',34);small=ImageFont.truetype('C:/Windows/Fonts/malgun.ttf',25)
 dr.text((im.width/2,im.height/2-22),label,font=f,fill=(65,77,89),anchor='mm')
 dr.text((im.width/2,im.height/2+30),'사진을 넣을 자리',font=small,fill=(100,110,120),anchor='mm')
 fp=work/f'slot{idx}.png';im.save(fp)
 x=p();x.paragraph_format.space_after=Pt(4);x.add_run().add_picture(str(fp),width=Inches(6.5))
 pic=doc.inline_shapes[-1];pic._inline.docPr.set('descr',label+' 증빙 사진 삽입 위치')
def link(label,url):
 x=p(); h=OxmlElement('w:hyperlink');h.set(qn('r:id'),doc.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True))
 r=OxmlElement('w:r');pr=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'245579');pr.append(c);r.append(pr);t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);x._p.append(h)
def record(name,stat,fields,height=1.7):
 doc.add_heading(name,level=2);p(stat);p(fields,9,'566674');slot(name,height)
# page 1
x=doc.add_paragraph('게임 플레이 경험',style='Title')
p('박신우  |  펄어비스 QA 인턴십 지원',11)
p('경쟁 게임, 생존 및 탐험 게임, 스포츠 게임을 플레이했습니다. 게임별 경험과 누적 플레이 시간, 랭크를 정리하고 관련 화면을 함께 제시합니다.')
heading('게임별 경험 요약')
rows=[['리그 오브 레전드','PC','에메랄드 달성','[기간 입력]'],['오버워치','PC 여부 확인','과거 다이아몬드 달성','[기간 입력]'],['Palworld 팰월드','Steam','289.1시간','[기간 입력]'],['Valheim','Steam','166.7시간','[기간 입력]'],['Rust','Steam','149.5시간','[기간 입력]'],['EA SPORTS FIFA 21','Steam','198.9시간','[기간 입력]'],['MLB The Show 25','Xbox 원격 이용','[시간 또는 진행도 입력]','[기간 입력]']]
t=doc.add_table(rows=1, cols=4);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
widths=[2.03,1.18,2.02,1.54]
for c,w,txt in zip(t.rows[0].cells,widths,['게임','플랫폼','플레이 기록','플레이 기간']):c.width=Inches(w);c.text=txt
for data in rows:
 for c,w,txt in zip(t.add_row().cells,widths,data):c.width=Inches(w);c.text=txt
for ri,row in enumerate(t.rows):
 trpr=row._tr.get_or_add_trPr(); ns=OxmlElement('w:cantSplit');trpr.append(ns)
 for c in row.cells:
  c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
  cp=c._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DFEAF1' if ri==0 else ('F5F7F9' if ri%2==0 else 'FFFFFF'));cp.append(sh)
  borders=OxmlElement('w:tcBorders')
  for side in ['top','left','bottom','right']:
   e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
  cp.append(borders);mar=OxmlElement('w:tcMar')
  for side in ['top','left','bottom','right']:
   e=OxmlElement('w:'+side);e.set(qn('w:w'),'100');e.set(qn('w:type'),'dxa');mar.append(e)
  cp.append(mar)
  for para in c.paragraphs:
   para.paragraph_format.space_after=Pt(2)
   for r in para.runs:r.font.size=Pt(9);r.bold=ri==0
p('누적 플레이 시간은 첨부한 Steam 화면 기준이며, 플레이 기간과는 별도로 표시합니다.',9,'566674')
link('Steam 게임 목록 및 플레이 시간 확인','https://steamcommunity.com/profiles/76561198149522185/games/?tab=all&sort=playtime')
doc.add_heading('대표 경험',level=2)
p('[가장 오래 즐겼거나 깊이 있게 플레이한 게임 2~3개와 주요 콘텐츠를 간단히 입력하세요.]',10,'566674')
# page 2
new('Steam 라이브러리 기록')
p('계정 tlsdn0403  |  게임명과 누적 플레이 시간이 표시된 화면')
src=Path('C:/Users/tlsdn/AppData/Local/Temp/codex-clipboard-48e2ec64-cf76-4184-bfa2-8d815e64a6b0.png')
if src.exists():
 p().add_run().add_picture(str(src),width=Inches(6.5));doc.inline_shapes[-1]._inline.docPr.set('descr','Steam 라이브러리 팰월드 FIFA 21 Valheim Rust 플레이 시간')
else:slot('Steam 라이브러리 첫 화면',3.6)
p('사진 기준일  [촬영일 입력]',9,'566674')
slot('추가 게임 목록 또는 최신 라이브러리 화면',2.25)
p('추가 설명  [다른 장르의 게임 또는 최근 플레이한 게임을 간단히 입력]',9,'566674')
# page 3
new('경쟁 게임 경험')
record('리그 오브 레전드','에메랄드 달성','기간 [입력]  |  달성 시즌 및 모드 [입력]  |  주 포지션 [입력]',2.0)
p('주요 경험  [주로 즐긴 콘텐츠와 기억에 남는 플레이 경험 1~2문장]',9,'566674')
record('오버워치','과거 다이아몬드 달성','기간 [입력]  |  달성 시즌 및 모드 [입력]  |  주 역할 및 영웅 [입력]',2.0)
p('주요 경험  [주로 즐긴 콘텐츠와 기억에 남는 플레이 경험 1~2문장]',9,'566674')
# page 4
new('생존과 탐험 게임 경험')
for name,hours in [('Palworld 팰월드','289.1'),('Valheim','166.7'),('Rust','149.5')]:
 record(name,'Steam 누적 '+hours+'시간','기간 [입력]  |  솔로 또는 멀티플레이 [입력]  |  주요 콘텐츠 및 진행도 [입력]',1.25)
 p('경험 메모  [건설·탐험·전투 등 실제 즐긴 내용 1문장]',9,'566674')
# page 5
new('스포츠 게임 경험')
record('EA SPORTS FIFA 21','Steam 누적 198.9시간','기간 [입력]  |  주로 즐긴 모드 [입력]  |  주요 기록 [입력]',2.0)
p('주요 경험  [실제로 즐긴 모드와 플레이 경험 1~2문장]',9,'566674')
record('MLB The Show 25','Xbox를 통한 원격 플레이 경험','기간 [입력]  |  이용 기기 [입력]  |  콘솔 원격 또는 클라우드 [확인 후 입력]',2.0)
p('주요 경험  [주로 즐긴 모드·진행도·타격 시스템 등 인상 깊었던 내용]',9,'566674')
# footer page number
foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=foot.add_run();fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');r._r.addnext(fld)
path=out/'박신우_게임플레이_경험_사진추가용.docx';doc.save(path)
(out/'게임플레이_문서_작성안내.txt').write_text('사진 추가 방법\n\n1. Word에서 DOCX를 엽니다.\n2. 회색 사진 자리를 오른쪽 클릭하고 그림 변경을 선택하여 본인 사진으로 교체합니다.\n3. 그림 변경이 없다면 해당 그림을 삭제하고 같은 위치에 삽입 > 그림으로 사진을 넣습니다. 배치는 글자처럼 취급 또는 텍스트 줄 안으로 설정합니다.\n4. 사진은 비율을 유지하고 폭 16.5cm 이내로 맞춥니다. 세로로 긴 화면은 잘라 여러 장으로 나누면 글씨가 읽기 쉽습니다.\n5. 대괄호 항목을 채우고 안내 문구를 삭제합니다. 기록이 없으면 확인 불가 또는 해당 없음으로 표시합니다.\n6. 사진이 없는 항목은 빈 사진 자리를 삭제해도 됩니다.\n7. 정확한 기간이 기억나지 않으면 약 또는 대략이라고 표시합니다. 스팀 시간은 최신 화면으로 갱신할 수 있습니다.\n8. 모두 채운 뒤 페이지가 밀리지 않았는지 확인하고 PDF로 저장해 제출합니다. 지원서의 파일 형식과 용량 제한을 확인하세요.\n\n현재 문서에는 기존 스팀 사진 한 장과 알려주신 기록이 반영되어 있습니다. 롤·오버워치 랭크 시즌, 각 게임 기간, Xbox 이용 방식은 확인 후 입력하세요.\n',encoding='utf-8')
print(path)
