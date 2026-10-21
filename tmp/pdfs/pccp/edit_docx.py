from docx import Document
from copy import deepcopy
from pathlib import Path
import shutil
p=Path('이력서/박신우_이력서.docx'); shutil.copy2(p,Path('tmp/pdfs/pccp/originals')/p.name)
d=Document(p)
heading=deepcopy(d.paragraphs[2]._p)
from docx.text.paragraph import Paragraph
hp=Paragraph(heading,d._body)
for run in hp.runs: run.text=''
hp.add_run('자격사항').bold=True
end=d.paragraphs[-1]._p; end.addprevious(heading)
table=deepcopy(d.tables[-1]._tbl); end.addprevious(table)
from docx.table import Table
t=Table(table,d._body)
vals=[['자격명','취득 등급','취득 여부','비고'],['PCCP','Level 3','취득','']]
for row,vs in zip(t.rows,vals):
    for c,s in zip(row.cells,vs):
        para=c.paragraphs[0]
        if para.runs:
            para.runs[0].text=s
            for r in para.runs[1:]:r.text=''
        else:para.add_run(s)
d.save(p)
print('DOCX updated')
