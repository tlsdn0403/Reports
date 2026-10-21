from pathlib import Path
from docx import Document
from docx.shared import Pt,Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from zipfile import ZipFile, ZIP_DEFLATED
from io import BytesIO
src=Path(r'C:/Users/tlsdn/OneDrive/문서/카카오톡 받은 파일/이력서 자소서 양식_260908.docx')
d=Document(src)
from copy import deepcopy
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
# Reuse the template's populated-row design for a third record in spacer slots.
for target,source in [(9,8),(24,23)]:
 table=d.tables[1]; old=table.rows[target]._tr; old.getparent().replace(old,deepcopy(table.rows[source]._tr))
def fill(t,r,c,text,size=9):
 cell=d.tables[t].rows[r].cells[c]
 p=cell.paragraphs[0]
 for other in list(cell.paragraphs)[1:]: other._element.getparent().remove(other._element)
 p.clear(); run=p.add_run(text); run.font.size=Pt(size)
 return p
for r,c,s in [(0,3,'박신우'),(0,5,'Park Shin Woo'),(0,7,'게임 클라이언트'),(1,3,'2001.04.03'),(1,7,'만 25세'),(3,3,'tlsdn0403@gmail.com'),(4,3,'경기도 용인시 수지구 정평로 73')]: fill(0,r,c,s)
p=fill(0,0,0,''); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p.add_run().add_picture('제출용 자료/증명사진.jpg',width=Cm(2.6))
for c,s in [(0,'2021.03~2027.02'),(3,'한국공학대학교 게임공학과'),(4,'3.5 / 4.5'),(6,'졸업 예정')]: fill(1,2,c,s)
for r,values in [(7,['2025.12~2026.07','감염: 죽음의 도시 · 절단 및 트럭 탑승','졸업작품','한국공학대학교']),(8,['2025.11~2025.12','NGP · 캐릭터 송수신 및 충돌 처리','전공 프로젝트','한국공학대학교']),(9,['2026.07~2026.08','FlickDom · 보드 판정 및 턴 동기화','팀 프로젝트','3인 팀'])]:
 for c,s in zip([0,3,4,6],values): fill(1,r,c,s)
for c,s in [(0,'OPIc'),(3,'IM2'),(4,'2023.10')]: fill(1,12,c,s)
for r,values in [(22,['Unreal Engine 5','C++ 게임플레이 구현','좀비 절단·트럭 탑승 구현']),(23,['Unity 6','C# 보드게임 개발','보드 판정·턴 상태 동기화']),(24,['OpenGL / GLSL','Shader Lab 렌더링 학습','광원·재질 조절 및 효과 비교'])]:
 for c,s in zip([0,3,5],values): fill(1,r,c,s)
fill(1,15,0,'자격 및 수상')
for c,text in [(0,'2026.10.07'),(1,'PCCP C++ Lv.3 · 700점'),(3,'주식회사 그렙')]:fill(1,17,c,text)
fill(1,18,1,'소프트웨어 경진대회 장려상 (PCCP 평가)')
fill(1,18,3,'한국공학대학교')
fill(1,26,4,'계급:'); fill(1,26,5,'전역(해제)사유: 군필')
answers=[
'졸업작품 「감염: 죽음의 도시」에서 UE5와 C++로 좀비 신체 절단과 트럭 탑승 기능을 맡았습니다. 트럭 위의 캐릭터가 튀어 오르거나 뒤로 밀리는 문제가 있었습니다. 먼저 차량과 캐릭터의 충돌을 분리하고, 트럭 바닥을 이동 기준으로 지정하는 SetBase를 적용해 차량을 따라가도록 했습니다. 이후 하차 시 위치가 순간적으로 보정되는 현상이 남아, SetBase(nullptr)로 이동 기준을 해제하고 관련 상태도 초기화했습니다. 이를 통해 충돌 처리와 함께 탑승·하차의 상태 전환도 확인해야 한다는 점을 배웠습니다.',
'수식이 화면에 어떤 변화를 만드는지 확인하며 공부합니다. 최근 C++17, OpenGL, GLSL을 사용한 Shader Lab을 AI와 협업해 구현했습니다. 광원의 위치와 세기, 재질 값을 실시간으로 조절하면서 Phong의 반사 벡터와 Blinn-Phong의 하프 벡터에 따라 하이라이트가 어떻게 달라지는지 비교했습니다. 소멸 효과에서는 노이즈와 임계값을 비교해 프래그먼트를 버리는 처리가 메시 자체를 삭제하는 것과 다름을 확인했습니다. LOD와 컬링 실험에서는 객체 수를 늘리며 삼각형 수, 드로 콜, GPU 시간을 함께 관찰했습니다. 코드를 수정한 뒤 셰이더를 다시 불러와 화면을 비교하는 방식으로 렌더링 원리를 학습하고 있습니다.',
'팀 프로젝트에서 선행 리팩터링이 늦어지면서 후속 작업을 맡은 팀원과 의견이 엇갈렸습니다. 소통까지 줄어들자 저는 두 팀원이 함께 이야기할 자리를 마련하고, 각자 어려운 점과 상대에게 바라는 점을 들었습니다. 누가 잘못했는지 판단하기보다 프로젝트를 함께 완성해야 한다는 목표를 중심으로 대화를 이어갔고, 이후 소통이 회복되어 작업을 계속할 수 있었습니다. 이 경험을 통해 의견이 다를 때 피드백을 미루기보다 상대의 상황을 먼저 듣고, 작업에 미치는 영향을 구체적으로 이야기해야 한다는 점을 배웠습니다.',
'C++로 게임에 필요한 데이터를 정의하고, 그 데이터를 캐릭터 동작에 반영할 수 있습니다. PCCP C++ Lv.3를 700점으로 취득했고, PCCP 점수로 평가한 한국공학대학교 소프트웨어 경진대회에서 장려상을 받았습니다. NGP 프로젝트에서는 다른 플레이어의 움직임을 동기화하기 위해 캐릭터 ID, 위치, 이동 방향, 회전각, 팔다리 움직임 각도를 전송 정보로 정했습니다. 이를 구조체에 담아 서버로 보내고, 수신한 정보를 다른 플레이어의 캐릭터 객체에 적용했습니다. 저는 캐릭터 정보 송수신과 적용, 충돌 처리를 맡았고 서버 연결과 스레드 구성은 팀원이 담당했습니다. Unity 프로젝트 FlickDom에서는 보드 점유와 카드 패턴 판정, 턴 상태 동기화를 구현하며 C# 활용 경험도 쌓았습니다.'
]
for r,s in zip([1,3,5,7],answers):
 p=fill(2,r,0,s,10); p.paragraph_format.line_spacing=Pt(14); p.paragraph_format.space_after=Pt(0)
 p.alignment=WD_ALIGN_PARAGRAPH.LEFT
 pr=p._p.get_or_add_pPr()
 for old in list(pr.findall(qn('w:snapToGrid'))):pr.remove(old)
 grid=OxmlElement('w:snapToGrid');grid.set(qn('w:val'),'0');pr.append(grid)
for p in d.paragraphs:
 if '성  명' in p.text or '성  명' in p.text.replace('\t',''):
  p.text='\t\t\t\t\t\t\t\t성  명 : 박신우      (인)'
# Recover vertical room from unused blank education/language rows.
for ri in [3,13]:
 h=d.tables[1].rows[ri]._tr.find(qn('w:trPr')).find(qn('w:trHeight'));h.set(qn('w:val'),'180')
buf=BytesIO(); d.save(buf); buf.seek(0)
out=Path('output/박신우_이력서_자기소개서_셰이더반영본.docx');out.parent.mkdir(exist_ok=True)
with ZipFile(src) as original,ZipFile(buf) as edited,ZipFile(out,'w',ZIP_DEFLATED) as final:
 for n in original.namelist():
  final.writestr(n, edited.read(n) if n in ['word/document.xml','word/_rels/document.xml.rels','[Content_Types].xml'] else original.read(n))
 for n in edited.namelist():
  if n not in original.namelist():final.writestr(n,edited.read(n))
print(out.resolve())
